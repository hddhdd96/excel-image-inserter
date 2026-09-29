"""
엑셀 파일 처리 클래스
"""

import os
import re
from typing import List, Optional, Tuple
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.drawing.image import Image as OpenpyxlImage
from PIL import Image as PILImage

from .error_codes import ErrorCode
from .logger import Logger


class ExcelHandler:
    """엑셀 파일 읽기/쓰기 처리"""
    
    def __init__(self, file_path: str, logger: Optional[Logger] = None):
        """
        Args:
            file_path: 엑셀 파일 경로
            logger: 로거 인스턴스
        """
        self.file_path = file_path
        self.logger = logger or Logger()
        self.workbook: Optional[Workbook] = None
        self.worksheet = None
        self.temp_files: List[str] = []  # 임시 파일 경로 리스트
    
    def load(self) -> bool:
        """
        엑셀 파일 로드
        
        Returns:
            성공 여부
        """
        try:
            if not os.path.exists(self.file_path):
                self.logger.error(f"{ErrorCode.E001}: 파일이 존재하지 않습니다: {self.file_path}")
                return False
            
            # data_only=True: 수식이 아닌 계산된 값을 가져옴
            self.workbook = load_workbook(self.file_path, data_only=True)
            self.logger.info(f"엑셀 파일 로드 완료: {self.file_path}")
            return True
        except Exception as e:
            self.logger.error(f"{ErrorCode.E001}: {str(e)}")
            return False
    
    def get_sheet_names(self) -> List[str]:
        """
        시트 이름 목록 반환
        
        Returns:
            시트 이름 리스트
        """
        if not self.workbook:
            return []
        return self.workbook.sheetnames
    
    def select_sheet(self, sheet_name: str) -> bool:
        """
        시트 선택
        
        Args:
            sheet_name: 시트 이름
            
        Returns:
            성공 여부
        """
        try:
            if not self.workbook:
                self.logger.error(f"{ErrorCode.E001}: 엑셀 파일이 로드되지 않았습니다")
                return False
            
            if sheet_name not in self.workbook.sheetnames:
                self.logger.error(f"시트를 찾을 수 없습니다: {sheet_name}")
                return False
            
            self.worksheet = self.workbook[sheet_name]
            self.logger.info(f"시트 선택: {sheet_name}")
            return True
        except Exception as e:
            self.logger.error(f"시트 선택 실패: {str(e)}")
            return False
    
    def get_column_index(self, column_letter: str) -> int:
        """
        열 문자를 인덱스로 변환 (A -> 1, B -> 2, ...)
        
        Args:
            column_letter: 열 문자 (예: 'A', 'B')
            
        Returns:
            열 인덱스 (1부터 시작)
        """
        try:
            return column_index_from_string(column_letter.upper())
        except Exception as e:
            self.logger.error(f"{ErrorCode.E003}: 열 좌표 변환 실패: {column_letter} - {str(e)}")
            return -1
    
    def parse_cell_coordinate(self, cell_ref: str) -> Tuple[int, int]:
        """
        셀 좌표 파싱 (예: 'B5' -> (5, 2))
        
        Args:
            cell_ref: 셀 참조 (예: 'B5')
            
        Returns:
            (행, 열) 튜플 (1부터 시작)
        """
        try:
            cell_ref = cell_ref.upper().strip()
            
            # 정규식으로 열과 행 분리 (예: "C5" -> 열: "C", 행: "5")
            # 열 부분: 하나 이상의 알파벳
            # 행 부분: 하나 이상의 숫자
            match = re.match(r'^([A-Z]+)(\d+)$', cell_ref)
            if not match:
                raise ValueError(f"잘못된 셀 참조 형식: {cell_ref}")
            
            col_letter = match.group(1)  # 예: "C"
            row_str = match.group(2)      # 예: "5"
            
            # 열을 인덱스로 변환
            col_idx = column_index_from_string(col_letter)
            row = int(row_str)
            
            return row, col_idx
        except Exception as e:
            self.logger.error(f"{ErrorCode.E003}: 셀 좌표 파싱 실패: {cell_ref} - {str(e)}")
            return -1, -1
    
    def read_column_values(self, column_letter: str, start_row: int = 1) -> List[Tuple[int, str]]:
        """
        열의 값들을 읽어서 반환 (빈 셀을 만나면 중단)
        
        Args:
            column_letter: 열 문자 (예: 'A')
            start_row: 시작 행 (1부터 시작)
            
        Returns:
            [(행, 값), ...] 리스트
        """
        if not self.worksheet:
            self.logger.error(f"{ErrorCode.E001}: 시트가 선택되지 않았습니다")
            return []
        
        column_idx = self.get_column_index(column_letter)
        if column_idx == -1:
            return []
        
        values = []
        row = start_row
        
        while True:
            cell = self.worksheet.cell(row=row, column=column_idx)
            cell_value = cell.value
            
            # 빈 셀이면 중단
            if cell_value is None or str(cell_value).strip() == "":
                break
            
            values.append((row, str(cell_value).strip()))
            row += 1
        
        return values
    
    def get_cell_dimensions(self, cell_ref: str) -> Tuple[float, float]:
        """
        셀의 크기(너비, 높이)를 픽셀 단위로 반환
        
        Args:
            cell_ref: 셀 참조 (예: 'B5')
            
        Returns:
            (너비, 높이) 튜플 (픽셀)
        """
        try:
            if not self.worksheet:
                return 0, 0
            
            row, col_idx = self.parse_cell_coordinate(cell_ref)
            if row == -1 or col_idx == -1:
                return 0, 0
            
            # 열 너비 가져오기 (문자 단위)
            col_letter = get_column_letter(col_idx)
            col_width = self.worksheet.column_dimensions[col_letter].width
            if col_width is None:
                col_width = 8.43  # 기본 열 너비
            
            # 행 높이 가져오기 (포인트 단위)
            row_height = self.worksheet.row_dimensions[row].height
            if row_height is None:
                row_height = 15  # 기본 행 높이
            
            # 픽셀로 변환
            # 열 너비: 1 문자 = 약 7 픽셀
            # 행 높이: 1 포인트 = 약 1.33 픽셀
            width_px = col_width * 7
            height_px = row_height * 1.33
            
            return width_px, height_px
        except Exception as e:
            self.logger.error(f"셀 크기 계산 실패 ({cell_ref}): {str(e)}")
            return 0, 0
    
    def insert_image(self, image_path: str, cell_ref: str, width: Optional[int] = None, height: Optional[int] = None, fit_to_cell: bool = False) -> bool:
        """
        이미지를 셀에 삽입
        
        Args:
            image_path: 이미지 파일 경로
            cell_ref: 셀 참조 (예: 'B5')
            width: 이미지 너비 (픽셀, None이면 원본 크기)
            height: 이미지 높이 (픽셀, None이면 원본 크기)
            fit_to_cell: True이면 셀 크기에 맞춤
            
        Returns:
            성공 여부
        """
        try:
            if not self.worksheet:
                self.logger.error(f"{ErrorCode.E001}: 시트가 선택되지 않았습니다")
                return False
            
            if not os.path.exists(image_path):
                self.logger.error(f"{ErrorCode.E004}: 이미지 파일이 존재하지 않습니다: {image_path}")
                return False
            
            # 이미지 로드
            img = PILImage.open(image_path)
            
            # 크기 결정
            target_width = None
            target_height = None
            
            if fit_to_cell:
                # 셀 크기에 맞춤
                cell_width, cell_height = self.get_cell_dimensions(cell_ref)
                if cell_width > 0 and cell_height > 0:
                    target_width = int(cell_width)
                    target_height = int(cell_height)
            elif width is not None or height is not None:
                # 수동 크기 설정
                target_width = width
                target_height = height
            
            # 크기 조정이 필요한 경우
            if target_width is not None or target_height is not None:
                original_width, original_height = img.size
                
                # 너비만 지정된 경우: 비율 유지하며 너비에 맞춤
                if target_width and not target_height:
                    ratio = target_width / original_width
                    target_height = int(original_height * ratio)
                # 높이만 지정된 경우: 비율 유지하며 높이에 맞춤
                elif target_height and not target_width:
                    ratio = target_height / original_height
                    target_width = int(original_width * ratio)
                
                # 실제 크기 조정 (resize 사용 - 정확한 크기로 조정)
                img = img.resize((target_width, target_height), PILImage.Resampling.LANCZOS)
            
            # 임시 파일로 저장 (openpyxl이 PIL Image를 직접 받지 못함)
            import tempfile
            with tempfile.NamedTemporaryFile(suffix='.png', delete=False) as tmp:
                img.save(tmp.name, 'PNG')
                tmp_path = tmp.name
            
            # openpyxl Image 객체 생성
            openpyxl_img = OpenpyxlImage(tmp_path)
            
            # 셀 위치에 이미지 삽입
            self.worksheet.add_image(openpyxl_img, cell_ref)
            
            # 임시 파일은 나중에 삭제하기 위해 리스트에 추가
            # (openpyxl은 save() 시점에 이미지 파일을 읽기 때문에 지금 삭제하면 안 됨)
            self.temp_files.append(tmp_path)
            
            return True
        except Exception as e:
            self.logger.error(f"{ErrorCode.E004}: 이미지 삽입 실패 ({cell_ref}): {str(e)}")
            return False
    
    def save(self, output_path: Optional[str] = None) -> bool:
        """
        엑셀 파일 저장
        
        Args:
            output_path: 저장 경로 (None이면 원본 파일에 덮어쓰기)
            
        Returns:
            성공 여부
        """
        try:
            if not self.workbook:
                self.logger.error(f"{ErrorCode.E001}: 엑셀 파일이 로드되지 않았습니다")
                return False
            
            save_path = output_path or self.file_path
            self.workbook.save(save_path)
            self.logger.info(f"엑셀 파일 저장 완료: {save_path}")
            
            # 저장 성공 후 임시 파일들 삭제
            self._cleanup_temp_files()
            
            return True
        except Exception as e:
            self.logger.error(f"엑셀 파일 저장 실패: {str(e)}")
            # 저장 실패해도 임시 파일은 정리
            self._cleanup_temp_files()
            return False
    
    def _cleanup_temp_files(self):
        """임시 파일들 정리"""
        for temp_path in self.temp_files:
            try:
                if os.path.exists(temp_path):
                    os.unlink(temp_path)
            except Exception as e:
                self.logger.warning(f"임시 파일 삭제 실패: {temp_path} - {str(e)}")
        self.temp_files.clear()
    
    def close(self):
        """엑셀 파일 닫기"""
        # 임시 파일 정리
        self._cleanup_temp_files()
        
        if self.workbook:
            self.workbook.close()
            self.workbook = None
            self.worksheet = None
