"""
이미지 삽입 처리 클래스
"""

import os
from typing import List, Tuple
from openpyxl.utils import get_column_letter

from .excel_handler import ExcelHandler
from .image_matcher import ImageMatcher
from .logger import Logger


class ImageProcessor:
    """이미지 삽입 처리"""
    
    def __init__(self, excel_handler: ExcelHandler, image_matcher: ImageMatcher, logger: Logger = None):
        """
        Args:
            excel_handler: 엑셀 핸들러
            image_matcher: 이미지 매처
            logger: 로거 인스턴스
        """
        self.excel_handler = excel_handler
        self.image_matcher = image_matcher
        self.logger = logger or Logger()
    
    def process(self, match_column: str, start_cell: str, fit_to_cell: bool = False, 
                width: int = None, height: int = None) -> Tuple[int, int, int, int, int, int]:
        """
        이미지 삽입 처리
        
        Args:
            match_column: 매칭 기준 열 (예: 'A')
            start_cell: 이미지 삽입 시작 셀 (예: 'B5')
            fit_to_cell: True이면 셀 크기에 맞춤
            width: 이미지 너비 (픽셀, fit_to_cell=False일 때만 사용)
            height: 이미지 높이 (픽셀, fit_to_cell=False일 때만 사용)
            
        Returns:
            (총 행 수, 행 성공 수, 행 실패 수, 총 이미지 수, 이미지 성공 수, 이미지 실패 수) 튜플
        """
        # 시작 셀 좌표 파싱
        start_row, start_col = self.excel_handler.parse_cell_coordinate(start_cell)
        if start_row == -1 or start_col == -1:
            self.logger.error(f"시작 셀 좌표 파싱 실패: {start_cell}")
            return 0, 0, 0, 0, 0, 0
        
        # 매칭 기준 열의 값들 읽기
        column_values = self.excel_handler.read_column_values(match_column, start_row=start_row)
        
        if not column_values:
            self.logger.warning("매칭할 데이터가 없습니다")
            return 0, 0, 0, 0, 0, 0
        
        total_rows = len(column_values)
        row_success_count = 0
        row_fail_count = 0
        
        # 이미지 통계 - 전체 로드된 이미지 수
        total_images = len(self.image_matcher.image_files)
        image_success_count = 0
        image_fail_count = 0
        
        self.logger.info(f"총 {total_rows}행 처리 시작")
        
        # 각 행 처리
        for row_idx, (row, cell_value) in enumerate(column_values):
            # 매칭된 이미지 찾기
            matching_images = self.image_matcher.find_matching_images(cell_value)
            
            if not matching_images:
                self.logger.fail(f"{match_column}{row} → 이미지 없음 (셀 값: {cell_value})")
                row_fail_count += 1
                continue
            
            # 셀 크기에 맞춤 모드일 경우, 첫 번째 셀의 크기를 미리 계산
            actual_width = width
            actual_height = height
            if fit_to_cell:
                # 첫 번째 이미지가 삽입될 셀의 크기 계산
                first_insert_col = start_col
                first_insert_col_letter = get_column_letter(first_insert_col)
                first_insert_cell = f"{first_insert_col_letter}{row}"
                cell_width, cell_height = self.excel_handler.get_cell_dimensions(first_insert_cell)
                if cell_width > 0 and cell_height > 0:
                    actual_width = int(cell_width)
                    actual_height = int(cell_height)
            
            # 이미지 삽입 (가로 방향)
            inserted_count = 0
            for img_idx, image_path in enumerate(matching_images):
                # 삽입할 셀 계산
                insert_col = start_col + img_idx
                insert_col_letter = get_column_letter(insert_col)
                insert_cell = f"{insert_col_letter}{row}"
                
                # 이미지 삽입 (크기 설정 포함, fit_to_cell=False로 고정하여 미리 계산한 크기 사용)
                if self.excel_handler.insert_image(image_path, insert_cell, 
                                                   width=actual_width, height=actual_height, 
                                                   fit_to_cell=False):
                    inserted_count += 1
                    image_success_count += 1
                    self.logger.success(f"{match_column}{row} → {os.path.basename(image_path)} 삽입 완료 ({insert_cell})")
                else:
                    image_fail_count += 1
                    self.logger.error(f"{match_column}{row} → {os.path.basename(image_path)} 삽입 실패")
            
            if inserted_count > 0:
                row_success_count += 1
            else:
                row_fail_count += 1
        
        # 이미지 실패 수 계산 (전체 이미지 - 성공한 이미지)
        image_fail_count = total_images - image_success_count
        
        # 요약
        self.logger.summary(f"[행 단위] 총 {total_rows}행 중 {row_success_count}개 성공, {row_fail_count}개 실패")
        self.logger.summary(f"[이미지 단위] 총 {total_images}개 중 {image_success_count}개 성공, {image_fail_count}개 실패")
        
        return total_rows, row_success_count, row_fail_count, total_images, image_success_count, image_fail_count
